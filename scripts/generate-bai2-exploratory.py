from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\TechStore")
OUTPUT = ROOT / "BAI_2_KIEM_THU_KHAM_PHA_PRODUCT_DETAIL.docx"
MOBILE_SCREENSHOT = Path(os.environ.get("TEMP", r"C:\Windows\Temp")) / "techstore_product_qa" / "mobile.png"

NAVY = "172554"
BLUE = "3158D4"
LIGHT_BLUE = "EAF0FF"
LIGHT_GRAY = "F4F6F8"
MID_GRAY = "D9DEE7"
GREEN = "DDF3E5"
GREEN_TEXT = "17633A"
RED = "FCE3E1"
RED_TEXT = "9E2B25"
AMBER = "FFF0CC"
AMBER_TEXT = "8A5A00"
WHITE = "FFFFFF"
BLACK = "111827"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        tag = "w:" + edge_name
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                element.set(qn("w:" + key), str(edge_data[key]))


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Times New Roman", size=11, bold=False, color=BLACK, italic=False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(p, text, **kwargs):
    r = p.add_run(text)
    set_run_font(r, **kwargs)
    return r


def set_paragraph(p, before=0, after=5, line=1.08, align=None, keep=False) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True


def clear_paragraph_borders(p) -> None:
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = p_bdr.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            p_bdr.append(node)
        node.set(qn("w:val"), "nil")


def add_bullet(doc, text, level=0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph(p, after=3, line=1.08)
    for run in p.runs:
        set_run_font(run, size=11)
    if not p.runs:
        add_text(p, text, size=11)
    else:
        p.runs[0].text = text


def add_numbered(doc, text) -> None:
    p = doc.add_paragraph(style="List Number")
    set_paragraph(p, after=3)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=11)
    else:
        add_text(p, text, size=11)


def add_heading(doc, text, level=1) -> None:
    p = doc.add_heading(text, level=level)
    set_paragraph(p, before=8 if level == 1 else 5, after=5, keep=True)
    size = 15 if level == 1 else 12
    color = NAVY if level == 1 else BLUE
    for r in p.runs:
        set_run_font(r, size=size, bold=True, color=color)
    return p


def add_info_box(doc, title: str, body: str, fill=LIGHT_BLUE, title_color=BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.55)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=160, end=160)
    set_cell_border(cell, top={"val": "single", "sz": "10", "color": title_color},
                    bottom={"val": "single", "sz": "4", "color": MID_GRAY},
                    start={"val": "single", "sz": "4", "color": MID_GRAY},
                    end={"val": "single", "sz": "4", "color": MID_GRAY})
    p = cell.paragraphs[0]
    set_paragraph(p, after=3)
    add_text(p, title, bold=True, color=title_color, size=11)
    p2 = cell.add_paragraph()
    set_paragraph(p2, after=0, line=1.08)
    add_text(p2, body, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def format_table(table, widths, font_size=9.2, header_fill=NAVY) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=85, bottom=85, start=90, end=90)
            set_cell_border(cell,
                            top={"val": "single", "sz": "4", "color": MID_GRAY},
                            bottom={"val": "single", "sz": "4", "color": MID_GRAY},
                            start={"val": "single", "sz": "4", "color": MID_GRAY},
                            end={"val": "single", "sz": "4", "color": MID_GRAY})
            if r_idx == 0:
                set_cell_shading(cell, header_fill)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            for p in cell.paragraphs:
                set_paragraph(p, after=0, line=1.02)
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(r_idx == 0), color=(WHITE if r_idx == 0 else BLACK))


def add_status(cell, text: str, kind: str) -> None:
    fill, color = {
        "pass": (GREEN, GREEN_TEXT),
        "fail": (RED, RED_TEXT),
        "partial": (AMBER, AMBER_TEXT),
        "blocked": (LIGHT_GRAY, "4B5563"),
    }[kind]
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.clear()
    add_text(p, text, size=8.8, bold=True, color=color)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "Trang ", size=9, color="64748B")
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color="64748B")


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Times New Roman"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

# Cover
p = doc.add_paragraph()
set_paragraph(p, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "BÀI 2", size=15, bold=True, color=BLUE)

p = doc.add_paragraph(style="Title")
set_paragraph(p, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
clear_paragraph_borders(p)
add_text(p, "KIỂM THỬ KHÁM PHÁ\nTRANG CHI TIẾT SẢN PHẨM", size=24, bold=True, color=BLACK)

p = doc.add_paragraph()
set_paragraph(p, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "Dự án website thương mại điện tử TechStore", size=13, bold=True, color=NAVY)

meta = doc.add_table(rows=5, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
meta_data = [
    ("Đối tượng kiểm thử", "Trang Product Detail Page của TechStore"),
    ("Sản phẩm mẫu", "Xiaomi Redmi Note 17 Pro Max 5G 12GB/256GB"),
    ("Timebox", "45 phút"),
    ("Vai trò", "Người kiểm thử khám phá"),
    ("Ngày thực hiện", "14/10/2026"),
]
for idx, (label, value) in enumerate(meta_data):
    meta.cell(idx, 0).text = label
    meta.cell(idx, 1).text = value
    set_cell_shading(meta.cell(idx, 0), LIGHT_BLUE)
    for c in meta.rows[idx].cells:
        set_cell_margins(c, top=105, bottom=105, start=120, end=120)
        set_cell_border(c,
                        top={"val": "single", "sz": "4", "color": MID_GRAY},
                        bottom={"val": "single", "sz": "4", "color": MID_GRAY},
                        start={"val": "single", "sz": "4", "color": MID_GRAY},
                        end={"val": "single", "sz": "4", "color": MID_GRAY})
        for p2 in c.paragraphs:
            set_paragraph(p2, after=0)
            for r in p2.runs:
                set_run_font(r, size=10.5, bold=(c == meta.cell(idx, 0)), color=(NAVY if c == meta.cell(idx, 0) else BLACK))
set_table_widths(meta, [1.65, 4.75])

doc.add_paragraph()
add_info_box(
    doc,
    "Kết luận nhanh",
    "Trang hiển thị đúng trên desktop và xử lý URL sản phẩm sai bằng trang 404. Tuy nhiên, giao diện ở chiều rộng 390 px bị tràn ngang và cắt nội dung quan trọng. Phiên kiểm thử ghi nhận 01 lỗi Major, 01 rủi ro Medium cần xác nhận E2E và 02 khoảng trống về yêu cầu/dữ liệu kiểm thử.",
    fill="EEF2FF",
)

p = doc.add_paragraph()
set_paragraph(p, before=10, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "Tài liệu trả lời trực tiếp hai nhiệm vụ: Test Charter và kết quả phiên Exploratory Testing 45 phút.", size=9.5, italic=True, color="64748B")

doc.add_page_break()

# 1. Charter
add_heading(doc, "1. Điều lệ kiểm thử (Test Charter)", 1)
add_heading(doc, "1.1. Mục tiêu chính", 2)
p = doc.add_paragraph()
set_paragraph(p, after=6, line=1.12)
add_text(p, "Trong 45 phút, khám phá và phát hiện các lỗi nghiêm trọng có thể làm khách hàng ", size=11)
add_text(p, "hiểu sai sản phẩm, không chọn được hàng hoặc không thể bắt đầu mua hàng", size=11, bold=True, color=NAVY)
add_text(p, ". Ưu tiên tính đúng của giá/tồn kho, ảnh và biến thể, số lượng, nút mua/giỏ hàng, khả năng hiển thị trên mobile và phản hồi khi dữ liệu đầu vào không hợp lệ.", size=11)

add_heading(doc, "1.2. Các khu vực chính cần tập trung", 2)
focus_rows = [
    ("Khu vực", "Nội dung khám phá", "Ưu tiên"),
    ("Điều hướng & nhận diện", "URL, breadcrumb, tên/brand/SKU, trạng thái 404 khi thiếu hoặc sai slug.", "Cao"),
    ("Ảnh sản phẩm", "Ảnh chính, thumbnail, ảnh lỗi, thay đổi ảnh theo biến thể.", "Cao"),
    ("Giá & khuyến giá", "Giá bán, giá gốc, phần trăm giảm; tính nhất quán giữa biến thể và tồn kho.", "Rất cao"),
    ("Biến thể & số lượng", "Chọn biến thể, trạng thái hết hàng, giới hạn 1..tồn kho, dữ liệu bị sửa ở client.", "Rất cao"),
    ("Mua hàng", "Thêm giỏ, Mua ngay, hành vi với khách chưa đăng nhập, quay lại đúng sản phẩm.", "Rất cao"),
    ("Nội dung hỗ trợ quyết định", "Mô tả, thông số, chính sách, sản phẩm liên quan, đánh giá sản phẩm.", "Trung bình"),
    ("UX & khả năng truy cập", "Responsive, bàn phím, nhãn nút, alt ảnh, thông báo trạng thái và lỗi.", "Cao"),
]
table = doc.add_table(rows=1, cols=3)
for j, h in enumerate(focus_rows[0]):
    table.cell(0, j).text = h
for row in focus_rows[1:]:
    cells = table.add_row().cells
    for j, v in enumerate(row):
        cells[j].text = v
format_table(table, [1.45, 4.15, 0.85], font_size=8.7)

add_heading(doc, "1.3. Những điều cần tìm hiểu trong phiên", 2)
for item in [
    "Giá, tồn kho và ảnh đang lấy theo sản phẩm hay theo biến thể; khi đổi biến thể, mọi trường có cập nhật đồng bộ không?",
    "Trang xử lý thế nào khi sản phẩm/biến thể không tồn tại, hết hàng hoặc số lượng vượt tồn kho?",
    "Khách chưa đăng nhập có quay lại đúng sản phẩm và tiếp tục hành động sau khi đăng nhập không?",
    "Trang có hoạt động ổn trên màn hình nhỏ, chỉ dùng bàn phím và khi ảnh không tải được không?",
    "Phần đánh giá sản phẩm có thuộc phạm vi chức năng dự kiến hay chưa được triển khai?",
]:
    add_bullet(doc, item)

add_heading(doc, "1.4. Phạm vi và tiêu chí đánh giá", 2)
p = doc.add_paragraph()
set_paragraph(p, after=3)
add_text(p, "Trong phạm vi: ", bold=True, color=NAVY)
add_text(p, "trang chi tiết sản phẩm, các điểm vào/ra trực tiếp của trang, phản hồi giao diện và HTTP liên quan.")
p = doc.add_paragraph()
set_paragraph(p, after=3)
add_text(p, "Ngoài phạm vi: ", bold=True, color=NAVY)
add_text(p, "kiểm thử đầy đủ giỏ hàng, checkout, thanh toán VNPay, quản trị, tải hiệu năng và pentest.")
p = doc.add_paragraph()
set_paragraph(p, after=3)
add_text(p, "Test oracle: ", bold=True, color=NAVY)
add_text(p, "tính nhất quán dữ liệu trên trang, mã phản hồi máy chủ, quy ước phổ biến của website bán hàng, khả năng hoàn thành nhiệm vụ và nguyên tắc responsive/accessibility cơ bản.")

doc.add_page_break()

# 2. Execution
add_heading(doc, "2. Thực hiện kiểm thử khám phá trong 45 phút", 1)
add_heading(doc, "2.1. Môi trường và dữ liệu", 2)
env_rows = [
    ("Hạng mục", "Thiết lập"),
    ("Ứng dụng", "TechStore chạy cục bộ tại http://localhost:8080/TechStore"),
    ("Trang mẫu", "/product?slug=xiaomi-redmi-note-17-pro-max-5g-12gb-256gb"),
    ("Trình duyệt", "Google Chrome; desktop 1440 × 1800 và mobile 390 × 1500"),
    ("Vai trò", "Khách chưa đăng nhập"),
    ("Dữ liệu mẫu", "Giá 12.590.000 đ; giá gốc 16.990.000 đ; giảm 26%; tồn kho 22"),
]
table = doc.add_table(rows=1, cols=2)
for j, h in enumerate(env_rows[0]):
    table.cell(0, j).text = h
for row in env_rows[1:]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = row
format_table(table, [1.5, 4.95], font_size=9.2)

add_heading(doc, "2.2. Phân bổ timebox", 2)
time_rows = [
    ("Thời gian", "Nội dung"),
    ("0–5 phút", "Lập mô hình nhanh của trang, xác định rủi ro và test oracle."),
    ("5–12 phút", "Điều hướng, dữ liệu cơ bản, URL hợp lệ/không hợp lệ và HTTP 404."),
    ("12–22 phút", "Ảnh, giá, giảm giá, SKU, biến thể, tồn kho và số lượng."),
    ("22–32 phút", "CTA Thêm giỏ/Mua ngay, hành vi khách chưa đăng nhập, kiểm tra biên."),
    ("32–39 phút", "Responsive desktop/mobile và kiểm tra accessibility nhanh."),
    ("39–43 phút", "Tour dữ liệu nhiều sản phẩm và nhận diện khoảng trống dữ liệu."),
    ("43–45 phút", "Tổng hợp phát hiện, mức độ nghiêm trọng và hướng kiểm thử tiếp."),
]
table = doc.add_table(rows=1, cols=2)
for j, h in enumerate(time_rows[0]):
    table.cell(0, j).text = h
for row in time_rows[1:]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = row
format_table(table, [1.25, 5.2], font_size=9.2)

add_heading(doc, "2.3. Nhật ký kiểm thử và kết quả", 2)
test_rows = [
    ("ID", "Thử nghiệm khám phá", "Kết quả quan sát", "Trạng thái"),
    ("ET-01", "Mở URL sản phẩm hợp lệ.", "HTTP 200; tên, giá, ảnh, tồn kho và CTA xuất hiện.", "Đạt"),
    ("ET-02", "Bỏ slug hoặc dùng slug không tồn tại.", "Máy chủ trả 404 đúng thay vì trang rỗng/500.", "Đạt"),
    ("ET-03", "Đối chiếu giá bán, giá gốc và % giảm.", "12,59 triệu so với 16,99 triệu tương ứng mức giảm làm tròn 26%.", "Đạt"),
    ("ET-04", "Kiểm tra ảnh chính và thumbnail.", "Ảnh tải được; chỉ có 01 ảnh nên chưa khám phá được chuyển nhiều ảnh.", "Một phần"),
    ("ET-05", "Đổi biến thể và quan sát giá/tồn kho/ảnh.", "Sản phẩm mẫu chỉ có 01 biến thể; không đủ dữ liệu để thực thi nhánh đổi biến thể.", "Bị chặn"),
    ("ET-06", "Thử biên số lượng 1 và tồn kho 22.", "UI vô hiệu nút giảm ở 1 và giới hạn theo stock; chưa giả mạo POST có đăng nhập.", "Một phần"),
    ("ET-07", "Bấm CTA khi chưa đăng nhập.", "CTA điều hướng đến đăng nhập; không phát sinh lỗi 500.", "Đạt"),
    ("ET-08", "Kiểm tra quay lại sản phẩm sau đăng nhập.", "Link đi thẳng /login; logic hiện tại có thể về /home nếu savedUrl chưa được lưu.", "Rủi ro"),
    ("ET-09", "Mở mô tả, thông số và sản phẩm liên quan.", "Các khối nội dung cần thiết đang hiển thị.", "Đạt"),
    ("ET-10", "Tìm phần đánh giá/xếp hạng.", "Không có khu vực đánh giá; cần xác nhận đây có phải yêu cầu hay không.", "Cần hỏi"),
    ("ET-11", "Quan sát desktop 1440 px.", "Bố cục 2 cột rõ, không thấy chồng lấn hoặc mất CTA.", "Đạt"),
    ("ET-12", "Quan sát mobile 390 px.", "Trang tràn ngang; header, tiêu đề, giá và panel mua bị cắt bên phải.", "Không đạt"),
    ("ET-13", "Kiểm tra nhanh semantic/keyboard qua markup.", "Có alt/aria-label ở phần chính; chưa hoàn tất toàn bộ tab order và screen reader.", "Một phần"),
    ("ET-14", "Tour 40 sản phẩm đang hoạt động.", "Mỗi sản phẩm mẫu chỉ có 01 ảnh và 01 biến thể; độ phủ nhánh tương tác thấp.", "Bị chặn"),
]
table = doc.add_table(rows=1, cols=4)
for j, h in enumerate(test_rows[0]):
    table.cell(0, j).text = h
status_kinds = {
    "Đạt": "pass", "Không đạt": "fail", "Một phần": "partial",
    "Bị chặn": "blocked", "Rủi ro": "partial", "Cần hỏi": "blocked"
}
for row in test_rows[1:]:
    cells = table.add_row().cells
    for j, v in enumerate(row):
        cells[j].text = v
    add_status(cells[3], row[3], status_kinds[row[3]])
format_table(table, [0.55, 1.75, 3.55, 0.75], font_size=7.9)

doc.add_page_break()

# Findings
add_heading(doc, "3. Phát hiện quan trọng", 1)

add_heading(doc, "F-01 — Trang bị tràn ngang ở mobile 390 px", 2)
finding = doc.add_table(rows=6, cols=2)
finding_data = [
    ("Mức độ", "Major / ưu tiên sửa trước khi phát hành mobile"),
    ("Bước tái hiện", "1) Mở trang sản phẩm; 2) đặt viewport 390 px; 3) quan sát header, breadcrumb, tiêu đề, giá và khối mua hàng."),
    ("Kết quả thực tế", "Nội dung vượt chiều rộng viewport; một phần header và thông tin mua hàng bị cắt, xuất hiện khả năng cuộn ngang."),
    ("Kết quả mong đợi", "Bố cục chuyển thành một cột; không cuộn ngang; giá và CTA đọc/thao tác được hoàn toàn."),
    ("Ảnh hưởng", "Người dùng điện thoại khó đọc giá hoặc bắt đầu mua hàng; ảnh hưởng trực tiếp đến chuyển đổi."),
    ("Gợi ý retest", "320, 360, 390, 412 và 768 px; kiểm tra min-width của grid/card và các phần tử header."),
]
for i, (label, value) in enumerate(finding_data):
    finding.cell(i, 0).text = label
    finding.cell(i, 1).text = value
    set_cell_shading(finding.cell(i, 0), RED if i == 0 else LIGHT_GRAY)
format_table(finding, [1.35, 5.1], font_size=9.2, header_fill=LIGHT_GRAY)
# Re-style all rows because this is a label/value table without header.
for i, row in enumerate(finding.rows):
    for j, cell in enumerate(row.cells):
        if j == 0:
            set_cell_shading(cell, RED if i == 0 else LIGHT_GRAY)
        else:
            set_cell_shading(cell, WHITE)
        for p2 in cell.paragraphs:
            for r in p2.runs:
                set_run_font(r, size=9.2, bold=(j == 0), color=(RED_TEXT if i == 0 and j == 0 else BLACK))

add_heading(doc, "F-02 — Rủi ro mất ngữ cảnh sản phẩm sau đăng nhập", 2)
add_info_box(
    doc,
    "Medium — cần xác nhận bằng kiểm thử E2E có tài khoản",
    "CTA của khách trỏ trực tiếp đến /login. Luồng đăng nhập chỉ quay lại savedUrl khi URL này đã được bộ lọc xác thực lưu trước đó; đi trực tiếp từ trang sản phẩm có thể không tạo savedUrl và khách bị đưa về /home. Mong đợi: sau đăng nhập, quay lại đúng sản phẩm hoặc tiếp tục hành động Thêm giỏ/Mua ngay.",
    fill=AMBER,
    title_color=AMBER_TEXT,
)

add_heading(doc, "F-03 — Câu hỏi về yêu cầu đánh giá sản phẩm", 2)
p = doc.add_paragraph()
set_paragraph(p, after=6)
add_text(p, "Không tìm thấy phần đánh giá/xếp hạng trên trang. Vì không có đặc tả chi tiết, đây là ", size=11)
add_text(p, "câu hỏi sản phẩm, chưa kết luận là lỗi", size=11, bold=True, color=NAVY)
add_text(p, ". Cần Product Owner xác nhận review/rating có nằm trong phạm vi phiên bản hiện tại hay không.", size=11)

add_heading(doc, "F-04 — Khoảng trống dữ liệu kiểm thử", 2)
p = doc.add_paragraph()
set_paragraph(p, after=6)
add_text(p, "Khi khảo sát 40 sản phẩm đang hoạt động, mỗi sản phẩm chỉ có 01 ảnh và 01 biến thể. Vì vậy chưa thể kiểm tra sâu: chuyển thumbnail, đổi màu/dung lượng, đồng bộ giá–ảnh–tồn kho và biến thể hết hàng. Cần bổ sung ít nhất một sản phẩm có 3 ảnh, 3 biến thể và 1 biến thể hết hàng.", size=11)

add_heading(doc, "4. Kết luận phiên kiểm thử", 1)
summary_rows = [
    ("Chỉ số", "Kết quả"),
    ("Kịch bản/ý tưởng đã khám phá", "14"),
    ("Lỗi xác nhận", "01 Major — responsive mobile"),
    ("Rủi ro cần xác nhận", "01 Medium — quay lại trang sau đăng nhập"),
    ("Câu hỏi/khoảng trống", "02 — đánh giá sản phẩm và dữ liệu nhiều ảnh/biến thể"),
    ("Khuyến nghị", "Chưa nên chốt giao diện mobile trước khi sửa F-01 và retest."),
]
table = doc.add_table(rows=1, cols=2)
for j, h in enumerate(summary_rows[0]):
    table.cell(0, j).text = h
for row in summary_rows[1:]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = row
format_table(table, [2.1, 4.35], font_size=9.4)

add_heading(doc, "4.1. Kiểm thử tiếp theo nếu có thêm thời gian", 2)
for item in [
    "Đăng nhập thật từ CTA sản phẩm, xác nhận URL quay lại và trạng thái giỏ hàng sau đăng nhập.",
    "Giả mạo quantity/variantId/CSRF ở request POST để xác nhận kiểm tra phía server.",
    "Bổ sung dữ liệu nhiều biến thể, nhiều ảnh và hết hàng rồi chạy lại ET-04, ET-05, ET-06.",
    "Kiểm tra keyboard-only, focus visible, screen reader và lỗi ảnh/CDN.",
    "Retest responsive trên thiết bị thật và các viewport 320–768 px sau khi sửa F-01.",
]:
    add_bullet(doc, item)

add_info_box(
    doc,
    "Phán quyết phiên 45 phút",
    "Trang đáp ứng luồng xem sản phẩm cơ bản trên desktop, nhưng còn lỗi UX nghiêm trọng trên mobile. Ưu tiên sửa F-01; sau đó xác nhận F-02 và bổ sung dữ liệu để mở rộng độ phủ cho ảnh/biến thể.",
    fill="EEF2FF",
)

# Optional evidence page.
if MOBILE_SCREENSHOT.exists():
    appendix_heading = add_heading(doc, "Phụ lục — Bằng chứng trực quan F-01", 1)
    p = doc.add_paragraph()
    set_paragraph(p, after=7)
    add_text(p, "Ảnh chụp viewport 390 × 1500 cho thấy nội dung phía phải bị cắt. Đây là bằng chứng của lỗi responsive được ghi nhận trong ET-12.", size=10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(MOBILE_SCREENSHOT), width=Inches(1.78))
    cap = doc.add_paragraph()
    set_paragraph(cap, before=3, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cap, "Hình 1. Trang chi tiết sản phẩm ở viewport 390 px — nội dung tràn/cắt bên phải.", size=9, italic=True, color="64748B")

# Core properties and save.
doc.core_properties.title = "Bài 2 - Kiểm thử khám phá trang chi tiết sản phẩm TechStore"
doc.core_properties.subject = "Exploratory Testing - Product Detail Page"
doc.core_properties.author = "Nhóm dự án TechStore"
doc.core_properties.keywords = "TechStore, exploratory testing, product detail page, test charter"
doc.save(OUTPUT)
print(OUTPUT)
